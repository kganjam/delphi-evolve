#!/usr/bin/env python3
"""Fallback slide file for organizers who want .pptx / Google Slides: one full-bleed 16:9 image per slide
(rendered from deck.html at 2x) with the speaker notes attached.  NOT editable text: the PDF and HTML are
the primary artifacts.  The organizers' accepted format is unknown (see planning/rethink/SUBMISSION.md).

    python3 presentation/export_pptx.py [presentation/deck.html]   ->  presentation/deck.pptx
"""
import re
import sys
import tempfile
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else HERE / "deck.html").resolve()
    notes = (src.parent / "speaker-notes.md").read_text()
    parts = re.split(r"\n\*\*Slide (\d)\*\* \([^)]*\)\n\n", notes)
    spoken = {int(parts[i]): parts[i + 1].strip() for i in range(1, len(parts), 2)}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 13.333 x 7.5 in
    with tempfile.TemporaryDirectory() as td, sync_playwright() as p:
        b = p.chromium.launch(args=["--disable-dev-shm-usage"])
        pg = b.new_context(viewport={"width": 1280, "height": 720}, device_scale_factor=2).new_page()
        for i in range(1, 7):
            pg.goto(src.as_uri() + f"#{i}")
            pg.reload()
            pg.wait_for_timeout(200)
            png = Path(td) / f"s{i}.png"
            pg.screenshot(path=str(png))
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            sl.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
            sl.notes_slide.notes_text_frame.text = spoken[i]
        b.close()
    out = src.with_suffix(".pptx")
    prs.save(str(out))
    print(f"{out} slides={len(prs.slides)}")


if __name__ == "__main__":
    main()
